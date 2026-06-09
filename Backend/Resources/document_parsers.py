import os
import zipfile
import xml.etree.ElementTree as ET
import csv
import base64
import requests
import hashlib
from typing import Dict, List, Any
from pypdf import PdfReader
from openpyxl import load_workbook
from pptx import Presentation

# Base64 image encoder
def encode_image_to_base64(image_path: str) -> str:
    with open(image_path, "rb") as image_file:
        return base64.b64encode(image_file.read()).decode('utf-8')

def get_mock_visual_insights(image_path: str) -> str:
    """Returns a neutral placeholder when no vision model is available. No static/fabricated content."""
    base_name = os.path.basename(image_path)
    return (
        f"[Image: {base_name}] "
        "Visual analysis not available - no vision model is currently installed. "
        "Install a vision-capable model (e.g. llama3.2-vision) in Ollama to enable automatic image analysis."
    )

# Calls Ollama Llama3.2-Vision to extract detailed text analysis of images
def get_image_insights(image_path: str, ollama_url: str = "http://127.0.0.1:11434") -> str:
    try:
        # Check if a local vision model is actually installed and available in Ollama
        has_vision_model = False
        try:
            tags_response = requests.get(f"{ollama_url}/api/tags", timeout=1)
            if tags_response.status_code == 200:
                installed_models = [m.get("name") for m in tags_response.json().get("models", [])]
                has_vision_model = any("vision" in m.lower() for m in installed_models)
        except Exception:
            pass

        if not has_vision_model:
            return get_mock_visual_insights(image_path)

        base64_image = encode_image_to_base64(image_path)
        prompt = (
            "Provide a highly detailed textual transcription and description of this project artifact. "
            "Identify all visible diagrams, Gantt chart segments, roadmaps, timelines, milestones, lists, "
            "metrics, action items, owner names, dates, and handwritten notes. "
            "Organize this information in a highly structured, readable plain text format."
        )
        
        try:
            # Call Ollama generate API
            response = requests.post(
                f"{ollama_url}/api/generate",
                json={
                    "model": "llama3.2-vision:latest",
                    "prompt": prompt,
                    "images": [base64_image],
                    "stream": False
                },
                timeout=10
            )
            if response.status_code == 200:
                return response.json().get("response", "No response returned from the vision model.")
            else:
                return get_mock_visual_insights(image_path)
        except Exception:
            return get_mock_visual_insights(image_path)
    except Exception as e:
        return f"Error executing image vision analysis: {str(e)}"

# Text file parser
def parse_txt(file_path: str) -> str:
    with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
        return f.read()

# Docx file parser (custom zip-XML parser for reliability)
def parse_docx(file_path: str) -> str:
    try:
        with zipfile.ZipFile(file_path) as docx:
            xml_content = docx.read('word/document.xml')
            root = ET.fromstring(xml_content)
            
            paragraphs = []
            for para in root.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p'):
                texts = [node.text for node in para.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t') if node.text]
                if texts:
                    paragraphs.append(''.join(texts))
            return '\n'.join(paragraphs)
    except Exception as e:
        return f"Error parsing Word Document: {str(e)}"

# CSV spreadsheet parser
def parse_csv(file_path: str) -> str:
    try:
        rows = []
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            reader = csv.reader(f)
            for row in reader:
                rows.append(' | '.join(row))
        return '\n'.join(rows)
    except Exception as e:
        return f"Error parsing CSV file: {str(e)}"

# XLSX Excel spreadsheet parser
def parse_xlsx(file_path: str) -> str:
    try:
        wb = load_workbook(file_path, read_only=True, data_only=True)
        sheets_data = []
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            sheets_data.append(f"--- Sheet: {sheet_name} ---")
            for row in sheet.iter_rows(values_only=True):
                # Filter out completely empty rows
                if any(cell is not None for cell in row):
                    row_str = ' | '.join(str(cell) if cell is not None else "" for cell in row)
                    sheets_data.append(row_str)
        return '\n'.join(sheets_data)
    except Exception as e:
        return f"Error parsing Excel spreadsheet: {str(e)}"

# PDF document parser
def parse_pdf(file_path: str) -> str:
    try:
        reader = PdfReader(file_path)
        pages_data = []
        for i, page in enumerate(reader.pages):
            text = page.extract_text()
            if text:
                pages_data.append(f"--- Page {i+1} ---\n{text}")
        return '\n'.join(pages_data)
    except Exception as e:
        return f"Error parsing PDF document: {str(e)}"

# PPTX PowerPoint parser
def parse_pptx(file_path: str) -> str:
    try:
        prs = Presentation(file_path)
        slides_data = []
        for i, slide in enumerate(prs.slides):
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        line = para.text.strip()
                        if line:
                            texts.append(line)
                if shape.has_table:
                    table = shape.table
                    for row in table.rows:
                        row_str = ' | '.join(cell.text.strip() for cell in row.cells)
                        if row_str.strip():
                            texts.append(row_str)
            if texts:
                slides_data.append(f"--- Slide {i+1} ---\n" + '\n'.join(texts))
        return '\n\n'.join(slides_data)
    except Exception as e:
        return f"Error parsing PowerPoint file: {str(e)}"

# Master router document parser
def parse_document(file_path: str, project_dir: str, ollama_url: str = "http://127.0.0.1:11434") -> str:
    _, ext = os.path.splitext(file_path.lower())
    
    if ext == '.txt':
        return parse_txt(file_path)
    elif ext == '.docx':
        return parse_docx(file_path)
    elif ext == '.csv':
        return parse_csv(file_path)
    elif ext == '.xlsx':
        return parse_xlsx(file_path)
    elif ext == '.pdf':
        return parse_pdf(file_path)
    elif ext == '.pptx':
        return parse_pptx(file_path)
    elif ext in ['.png', '.jpg', '.jpeg', '.bmp']:
        # Establish cache directory inside the specific project folder to cache the vision outputs
        cache_dir = os.path.join(project_dir, ".image_cache")
        os.makedirs(cache_dir, exist_ok=True)
        
        # Calculate image hash for caching
        file_hash = ""
        try:
            with open(file_path, 'rb') as f:
                file_hash = hashlib.md5(f.read()).hexdigest()
        except:
            file_hash = os.path.basename(file_path)
            
        cache_file_name = f"{file_hash}_insight.txt"
        cache_file_path = os.path.join(cache_dir, cache_file_name)
        
        # Check if visual insights have already been computed and cached
        if os.path.exists(cache_file_path):
            with open(cache_file_path, 'r', encoding='utf-8') as f:
                return f.read()
                
        # Generate new visual insights from Ollama vision model
        insights = get_image_insights(file_path, ollama_url=ollama_url)
        
        # Write description insights to cache file
        with open(cache_file_path, 'w', encoding='utf-8') as f:
            f.write(insights)
            
        return insights
    else:
        return f"Unsupported file format: {ext}"
